"""
Isolates the upload format. If the intake sheet's columns ever shift, this is
the one file that should need to change.
"""
import io

import openpyxl

from backend.app.models import DeterminationResult, PriorAuthRequest

REQUIRED_COLUMNS = [
    "request_id",
    "member_id",
    "age",
    "procedure",
    "clinical_note",
    "requesting_provider",
]


def parse_requests(file_bytes: bytes) -> list[PriorAuthRequest]:
    """Read the nurse's uploaded xlsx into typed request rows."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() for h in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"Uploaded sheet is missing required columns: {missing}")

    col_index = {name: header.index(name) for name in REQUIRED_COLUMNS}
    requests = []
    for row in rows[1:]:
        if row[col_index["request_id"]] is None:
            continue
        requests.append(
            PriorAuthRequest(
                request_id=row[col_index["request_id"]],
                member_id=row[col_index["member_id"]],
                age=row[col_index["age"]],
                procedure=row[col_index["procedure"]],
                clinical_note=row[col_index["clinical_note"]],
                requesting_provider=row[col_index["requesting_provider"]],
            )
        )
    return requests


def write_results_xlsx(results: list[DeterminationResult]) -> bytes:
    """Mirror of parse_requests for the download side of the round trip."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Determinations"
    columns = [
        "request_id",
        "member_id",
        "procedure",
        "determination",
        "policy_id",
        "criterion",
        "reason",
        "urgent",
    ]
    ws.append(columns)
    for r in results:
        ws.append([getattr(r, c) for c in columns])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
